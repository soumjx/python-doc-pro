import os
import logging
from pathlib import Path
from typing import List, Optional

import mammoth
from xhtml2pdf import pisa
import pandas as pd
from docx import Document
from markdownify import markdownify as md

# Configure Logger
logger = logging.getLogger(__name__)

class DocxConverter:
    """
    A professional wrapper to convert DOCX files to HTML, MD, PDF, and CSV.
    """

    def __init__(self, input_path: str, output_dir: str):
        self.input_path = Path(input_path)
        self.output_dir = Path(output_dir)
        self.filename = self.input_path.stem

        if not self.input_path.exists():
            raise FileNotFoundError(f"Input file not found: {self.input_path}")
        
        # Ensure output directory exists
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def convert_to_html(self) -> str:
        """Converts DOCX to HTML using Mammoth for semantic accuracy."""
        try:
            output_file = self.output_dir / f"{self.filename}.html"
            logger.info(f"Converting to HTML: {output_file}")

            with open(self.input_path, "rb") as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html = result.value
                messages = result.messages # Warnings

            # Wrap in basic HTML structure for valid standalone files
            full_html = f"<html><body>{html}</body></html>"
            
            with open(output_file, "w", encoding="utf-8") as f:
                f.write(full_html)
            
            return full_html
        except Exception as e:
            logger.error(f"Failed to convert HTML: {e}")
            raise

    def convert_to_markdown(self) -> None:
        """Converts DOCX to Markdown (via HTML intermediate)."""
        try:
            output_file = self.output_dir / f"{self.filename}.md"
            logger.info(f"Converting to Markdown: {output_file}")

            # Reuse HTML conversion logic
            with open(self.input_path, "rb") as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html = result.value

            markdown_text = md(html, heading_style="ATX")

            with open(output_file, "w", encoding="utf-8") as f:
                f.write(markdown_text)

        except Exception as e:
            logger.error(f"Failed to convert Markdown: {e}")

    def convert_to_csv(self) -> None:
        """
        Extracts tables from DOCX and saves them as individual CSVs.
        Note: Word docs are not flat data, so we extract tables specifically.
        """
        try:
            logger.info("Extracting tables to CSV...")
            doc = Document(self.input_path)
            
            if not doc.tables:
                logger.warning("No tables found in document to convert to CSV.")
                return

            for i, table in enumerate(doc.tables):
                data = []
                keys = None
                
                for j, row in enumerate(table.rows):
                    text = (cell.text for cell in row.cells)
                    if j == 0:
                        keys = tuple(text)
                    else:
                        row_data = dict(zip(keys, text))
                        data.append(row_data)
                
                if data:
                    df = pd.DataFrame(data)
                    csv_filename = self.output_dir / f"{self.filename}_table_{i+1}.csv"
                    df.to_csv(csv_filename, index=False)
                    logger.info(f"Saved table {i+1} to {csv_filename}")

        except Exception as e:
            logger.error(f"Failed to extract CSV: {e}")

    def convert_to_pdf(self) -> None:
        """
        Converts DOCX to PDF using xhtml2pdf (via HTML intermediate).
        Pure Python solution, no external binaries required.
        """
        try:
            output_file = self.output_dir / f"{self.filename}.pdf"
            logger.info(f"Converting to PDF: {output_file}")

            # Get HTML content first
            with open(self.input_path, "rb") as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html = result.value

            full_html = f"<html><body>{html}</body></html>"
            
            # Convert using xhtml2pdf
            with open(output_file, "wb") as pdf_file:
                pisa_status = pisa.CreatePDF(full_html, dest=pdf_file)
                
            if pisa_status.err:
                logger.error(f"PDF generation error: {pisa_status.err}")
            else:
                logger.info("PDF generated successfully.")

        except Exception as e:
            logger.error(f"Failed to convert PDF: {e}")

    def process_all(self):
        """Runs all conversions."""
        self.convert_to_html()
        self.convert_to_markdown()
        self.convert_to_csv()
        self.convert_to_pdf()