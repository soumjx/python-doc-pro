from docx import Document

def create_table_docx(filename):
    doc = Document()
    doc.add_heading('Table Test Document', 0)

    doc.add_paragraph('This document contains a table for testing CSV extraction.')

    # Create a table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Name'
    hdr_cells[2].text = 'Role'

    # Add data rows
    data = [
        ('1', 'Alice', 'Engineer'),
        ('2', 'Bob', 'Designer'),
        ('3', 'Charlie', 'Manager')
    ]

    for id_val, name, role in data:
        row_cells = table.add_row().cells
        row_cells[0].text = id_val
        row_cells[1].text = name
        row_cells[2].text = role

    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_table_docx("input/tables.docx")
